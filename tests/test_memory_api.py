"""CP3B-1: memory persistence + minimal API integration tests.

Schema auto-applies on Database init; endpoints and lifecycle run against
real Indexer/DB paths with a stubbed model client (no network, no spend).
"""
import os
import tempfile
from pathlib import Path

import pytest
from fastapi import HTTPException

os.environ.setdefault("YAAD_DB_PATH", str(Path(tempfile.gettempdir()) / "yaad-memory-api-tests.sqlite3"))

from yaad.api import main
from yaad.config import Settings
from yaad.db import Database
import yaad.llm.meta_client as meta_client
from yaad.memory.extract import run_extraction
from yaad.scan.worker import Indexer


class StubClient:
    def __init__(self, records):
        self.records = records
        self.last_call = {"provider": "Meta Model API", "model": "muse-spark-1.3",
                          "status_code": 200, "success": True, "duration_ms": 5,
                          "input_tokens": 900, "output_tokens": 300, "cached_tokens": 0}

    def reason_json(self, messages, schema, prompt_id):
        return {"records": self.records}


def candidate(**over):
    base = {"record_type": "deadline", "title": "Draft due", "detail": "Draft due March 3rd.",
            "date_value": "2026-03-03", "date_raw": "March 3rd", "assignee": "",
            "status": "open", "statement_kind": "explicit", "key_phrase": "deliver draft",
            "evidence": [{"evidence_id": "PLACEHOLDER", "role": "supports",
                          "excerpt": "Deliver the draft by March 3rd"}]}
    base.update(over)
    return base


@pytest.fixture
def store(tmp_path):
    database = Database(tmp_path / "mem.sqlite3")
    folder = tmp_path / "src"
    folder.mkdir()
    settings = Settings(db_path=tmp_path / "mem.sqlite3", cloud_enabled=True,
                        model_api_key="stub-key-for-tests")
    indexer = Indexer(database, Settings(db_path=tmp_path / "mem.sqlite3",
                                         cloud_enabled=False))
    folder_id = indexer.add_folder(str(folder), True)
    return database, indexer, folder, folder_id, settings


def indexed_evidence(database, folder_id):
    with database.connect() as con:
        return [dict(r) for r in con.execute(
            "SELECT u.evidence_id,u.text FROM units u JOIN files f ON f.id=u.file_id "
            "WHERE f.folder_id=? ORDER BY u.id", (folder_id,))]


def test_schema_applies_on_fresh_database_init(tmp_path):
    database = Database(tmp_path / "fresh.sqlite3")
    with database.connect() as con:
        tables = {r["name"] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"memory_records", "memory_record_evidence", "record_flags"} <= tables


def test_extract_endpoint_stores_records_and_reports_rejected(store, monkeypatch):
    database, indexer, folder, folder_id, settings = store
    (folder / "plan.txt").write_text("Deliver the draft by March 3rd.", encoding="utf-8")
    indexer.index_folder(folder_id)
    eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    good = candidate(evidence=[{"evidence_id": eid, "role": "supports",
                                "excerpt": "Deliver the draft by March 3rd"}])
    bad = candidate(title="Bad date", date_value="someday")
    monkeypatch.setattr(main, "db", database)
    monkeypatch.setattr(main, "settings", settings)
    monkeypatch.setattr(meta_client, "MetaClient", lambda _s: StubClient([good, bad]))
    summary = main.extract_memory(main.MemoryExtractRequest(folder_id=folder_id))
    assert summary["accepted"] == 1 and summary["upserted"] == 1
    assert len(summary["rejected"]) == 1 and summary["rejected"][0]["errors"]
    assert summary["folder_id"] == folder_id and summary["spend_usd"] > 0
    with database.connect() as con:
        assert con.execute("SELECT COUNT(*) n FROM memory_records").fetchone()["n"] == 1
        assert con.execute("SELECT purpose FROM egress_log").fetchone()["purpose"] == "memory extraction"


def test_extract_endpoint_errors(store, monkeypatch, tmp_path):
    database, indexer, folder, folder_id, settings = store
    monkeypatch.setattr(main, "db", database)
    monkeypatch.setattr(main, "settings", settings)
    monkeypatch.setattr(meta_client, "MetaClient", lambda _s: StubClient([]))
    with pytest.raises(HTTPException) as missing:
        main.extract_memory(main.MemoryExtractRequest(folder_id=9999))
    assert missing.value.status_code == 404
    monkeypatch.setattr(main, "settings", Settings(db_path=tmp_path / "mem.sqlite3",
                                                   cloud_enabled=False))
    with pytest.raises(HTTPException) as gated:
        main.extract_memory(main.MemoryExtractRequest(folder_id=folder_id))
    assert gated.value.status_code == 409
    assert main._memory_lock.acquire(blocking=False)
    try:
        with pytest.raises(HTTPException) as busy:
            main.extract_memory(main.MemoryExtractRequest(folder_id=folder_id))
        assert busy.value.status_code == 409
    finally:
        main._memory_lock.release()


def test_list_endpoint_reports_evidence_and_serving_verdicts(store, monkeypatch):
    database, indexer, folder, folder_id, settings = store
    (folder / "plan.txt").write_text("Deliver the draft by March 3rd. Perhaps Maya could review it.",
                                     encoding="utf-8")
    indexer.index_folder(folder_id)
    eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    explicit = candidate(evidence=[{"evidence_id": eid, "role": "supports",
                                    "excerpt": "Deliver the draft by March 3rd"}])
    inferred = candidate(record_type="decision", title="Maya reviews", detail="Maya may review.",
                         date_value="", date_raw="", status="decided", statement_kind="inferred",
                         key_phrase="maya review",
                         evidence=[{"evidence_id": eid, "role": "supports",
                                    "excerpt": "Perhaps Maya could review it"}])
    run_extraction(database, folder_id, "run-list", settings, StubClient([explicit, inferred]))
    monkeypatch.setattr(main, "db", database)
    payload = main.list_memory(folder_id=folder_id)
    assert len(payload["records"]) == 2
    by_title = {r["title"]: r for r in payload["records"]}
    assert by_title["Draft due"]["serving"]["verdict"] == "eligible"
    assert by_title["Draft due"]["evidence"][0]["file"].endswith("plan.txt")
    assert by_title["Draft due"]["resolving_supports"] == 1
    # Inferred records are listed for inspection but never eligible as facts.
    assert by_title["Maya reviews"]["serving"]["verdict"] == "excluded"
    with pytest.raises(HTTPException) as missing:
        main.list_memory(folder_id=9999)
    assert missing.value.status_code == 404


def test_flag_endpoint_records_correction_requests(store, monkeypatch):
    database, indexer, folder, folder_id, settings = store
    (folder / "plan.txt").write_text("Deliver the draft by March 3rd.", encoding="utf-8")
    indexer.index_folder(folder_id)
    eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    run_extraction(database, folder_id, "run-flag", settings, StubClient(
        [candidate(evidence=[{"evidence_id": eid, "role": "supports",
                              "excerpt": "Deliver the draft by March 3rd"}])]))
    monkeypatch.setattr(main, "db", database)
    with database.connect() as con:
        record_id = con.execute("SELECT id FROM memory_records").fetchone()["id"]
    result = main.flag_memory(record_id, main.MemoryFlagRequest(reason="wrong date"))
    assert result["record_id"] == record_id
    with database.connect() as con:
        assert con.execute("SELECT reason FROM record_flags").fetchone()["reason"] == "wrong date"
    with pytest.raises(HTTPException) as missing:
        main.flag_memory(9999, main.MemoryFlagRequest(reason="x"))
    assert missing.value.status_code == 404


def test_repeated_extraction_is_idempotent(store):
    database, indexer, folder, folder_id, settings = store
    (folder / "plan.txt").write_text("Deliver the draft by March 3rd.", encoding="utf-8")
    indexer.index_folder(folder_id)
    eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    cand = candidate(evidence=[{"evidence_id": eid, "role": "supports",
                                "excerpt": "Deliver the draft by March 3rd"}])
    first = run_extraction(database, folder_id, "run-1", settings, StubClient([cand]))
    second = run_extraction(database, folder_id, "run-2", settings, StubClient([cand]))
    assert (first["accepted"], second["accepted"]) == (1, 1)
    with database.connect() as con:
        rows = con.execute("SELECT id,record_key,extraction_run FROM memory_records").fetchall()
        assert len(rows) == 1 and rows[0]["extraction_run"] == "run-2"
        assert con.execute("SELECT COUNT(*) n FROM memory_record_evidence").fetchone()["n"] == 1


def test_reindex_marks_stale_and_next_extraction_restores(store):
    database, indexer, folder, folder_id, settings = store
    target = folder / "plan.txt"
    target.write_text("Deliver the draft by March 3rd.", encoding="utf-8")
    indexer.index_folder(folder_id)
    old_eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    run_extraction(database, folder_id, "run-1", settings, StubClient(
        [candidate(evidence=[{"evidence_id": old_eid, "role": "supports",
                              "excerpt": "Deliver the draft by March 3rd"}])]))
    target.write_text("Deliver the draft by March 3rd. Added context paragraph.", encoding="utf-8")
    indexer.index_folder(folder_id)
    with database.connect() as con:
        assert con.execute("SELECT status FROM memory_records").fetchone()["status"] == "stale"
    new_eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    assert new_eid != old_eid
    run_extraction(database, folder_id, "run-2", settings, StubClient(
        [candidate(evidence=[{"evidence_id": new_eid, "role": "supports",
                              "excerpt": "Deliver the draft by March 3rd"}])]))
    with database.connect() as con:
        row = con.execute("SELECT status,extraction_run FROM memory_records").fetchone()
        assert (row["status"], row["extraction_run"]) == ("open", "run-2")
        assert con.execute("SELECT evidence_id FROM memory_record_evidence").fetchone()[0] == new_eid


def test_deleted_source_marks_record_stale(store):
    database, indexer, folder, folder_id, settings = store
    target = folder / "plan.txt"
    target.write_text("Deliver the draft by March 3rd.", encoding="utf-8")
    indexer.index_folder(folder_id)
    eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    run_extraction(database, folder_id, "run-1", settings, StubClient(
        [candidate(evidence=[{"evidence_id": eid, "role": "supports",
                              "excerpt": "Deliver the draft by March 3rd"}])]))
    target.unlink()
    indexer.index_folder(folder_id)
    with database.connect() as con:
        assert con.execute("SELECT COUNT(*) n FROM units").fetchone()["n"] == 0
        assert con.execute("SELECT status FROM memory_records").fetchone()["status"] == "stale"


def test_changed_date_claim_conflicts_through_service(store):
    database, indexer, folder, folder_id, settings = store
    (folder / "plan.txt").write_text("Deliver the draft by March 3rd. Update: draft moved to March 10th instead.",
                                     encoding="utf-8")
    indexer.index_folder(folder_id)
    eid = indexed_evidence(database, folder_id)[0]["evidence_id"]
    first = candidate(key_phrase="march draft",
                      evidence=[{"evidence_id": eid, "role": "supports",
                                 "excerpt": "Deliver the draft by March 3rd"}])
    moved = candidate(title="Draft due moved", detail="Draft moved to March 10th.",
                      date_value="2026-03-10", date_raw="March 10th", key_phrase="march draft",
                      evidence=[{"evidence_id": eid, "role": "supports",
                                 "excerpt": "draft moved to March 10th"}])
    run_extraction(database, folder_id, "run-1", settings, StubClient([first]))
    run_extraction(database, folder_id, "run-2", settings, StubClient([moved]))
    with database.connect() as con:
        assert {r["status"] for r in con.execute("SELECT status FROM memory_records")} == {"conflicted"}


def test_supersession_steps_down_older_claim_through_service(store):
    from docx import Document

    database, indexer, folder, folder_id, settings = store
    doc = Document()
    doc.add_paragraph("The venue is the old hall.")
    doc.add_paragraph("Correction: the venue is the new auditorium.")
    doc.save(folder / "venue.docx")
    indexer.index_folder(folder_id)
    old_eid, new_eid = [u["evidence_id"] for u in indexed_evidence(database, folder_id)]
    old = candidate(record_type="decision", title="Venue old hall", detail="Venue is the old hall.",
                    date_value="", date_raw="", status="decided", key_phrase="venue hall",
                    evidence=[{"evidence_id": old_eid, "role": "supports",
                               "excerpt": "The venue is the old hall"}])
    new = candidate(record_type="decision", title="Venue auditorium", detail="Venue is the new auditorium.",
                    date_value="", date_raw="", status="decided", key_phrase="venue auditorium",
                    evidence=[{"evidence_id": new_eid, "role": "supports",
                               "excerpt": "the venue is the new auditorium"},
                              {"evidence_id": old_eid, "role": "supersedes",
                               "excerpt": "The venue is the old hall"}])
    run_extraction(database, folder_id, "run-1", settings, StubClient([old]))
    run_extraction(database, folder_id, "run-2", settings, StubClient([new]))
    with database.connect() as con:
        by_title = {r["title"]: r["status"] for r in con.execute("SELECT title,status FROM memory_records")}
    assert by_title == {"Venue old hall": "superseded", "Venue auditorium": "decided"}
